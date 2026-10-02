# -*- coding: utf-8 -*-
"""/api/davomat/... — JSON-эндпоинты отчётов, check-in и администрирования."""
from __future__ import annotations

import html
import io
import json
import math
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from pydantic import BaseModel
from sqlalchemy import and_, select

import telegram
from events import _format_time_uz
from config import BOT_USERNAME, CHECKIN_PHOTO_DIR, COOKIE_SECURE, DEV_LOGIN, SESSION_TTL_HOURS
from database import (
    app_users, audit_log, checkins, corrections, database, employee_roles, employees,
    resolutions, roles, settings, tg_link_requests, web_credentials,
)
from davomat import service
from davomat.auth import (
    ADMIN_LEVELS, COOKIE_NAME, SUPERADMIN, VIEW_LEVELS, AuthError, bot_token, current_actor, make_token,
    require_admin, require_superadmin,
    hash_password, require_employee, require_viewer, verify_login_widget, verify_password,
    verify_webapp_init_data,
)
from davomat.engine import RoleParams, SHIFT_GUARD_DAY, SHIFT_GUARD_NIGHT, fmt_duration

router = APIRouter(prefix="/api/davomat")
PHOTO_DIR = Path(CHECKIN_PHOTO_DIR)
MAX_RANGE_DAYS = 366


def _d(s: Optional[str], default: date) -> date:
    if not s:
        return default
    try:
        return date.fromisoformat(s)
    except ValueError:
        raise HTTPException(400, f"Sana noto'g'ri: {s}")


def _range(frm: Optional[str], to: Optional[str], default_days: int = 7) -> tuple[date, date]:
    today = service.now_local().date()
    d_to = _d(to, today)
    d_from = _d(frm, d_to - timedelta(days=default_days - 1))
    if d_from > d_to:
        raise HTTPException(400, "Boshlanish sanasi tugashdan keyin")
    if (d_to - d_from).days >= MAX_RANGE_DAYS:
        raise HTTPException(400, f"Oraliq {MAX_RANGE_DAYS} kundan (1 yil) oshmasin")
    return d_from, d_to


def _session_response(uid: int, name: str, level, employee_no) -> JSONResponse:
    try:
        token = make_token(uid, name)
    except AuthError as e:
        raise HTTPException(500, str(e))
    resp = JSONResponse({"token": token, "uid": uid, "name": name, "level": level, "employee_no": employee_no})
    resp.set_cookie(COOKIE_NAME, token, max_age=SESSION_TTL_HOURS * 3600, httponly=True,
                    secure=COOKIE_SECURE, samesite="lax", path="/")
    return resp


async def _login(uid: int, name: str) -> JSONResponse:
    u = await database.fetch_one(select(app_users).where(app_users.c.telegram_user_id == uid))
    e = await database.fetch_one(select(employees).where(employees.c.telegram_user_id == uid))
    if not u and not e:
        raise HTTPException(403, f"Ruxsat yo'q. Admin sizni qo'shishi kerak (Telegram ID: {uid})")
    return _session_response(uid, name, u["level"] if u else None, e["employee_no"] if e else None)


# ------------------------------- авторизация -------------------------------

class WebAppAuth(BaseModel):
    init_data: str


@router.get("/config")
async def public_config():
    return {"bot_username": BOT_USERNAME, "dev_login": DEV_LOGIN}


@router.post("/auth/webapp")
async def auth_webapp(body: WebAppAuth):
    try:
        user = verify_webapp_init_data(body.init_data, bot_token())
    except AuthError as e:
        raise HTTPException(401, f"Telegram tekshiruvi o'tmadi: {e}")
    return await _login(int(user["id"]), " ".join(filter(None, [user.get("first_name"), user.get("last_name")])))


@router.post("/auth/widget")
async def auth_widget(request: Request):
    data = await request.json()
    try:
        user = verify_login_widget(data, bot_token())
    except AuthError as e:
        raise HTTPException(401, f"Telegram tekshiruvi o'tmadi: {e}")
    return await _login(user["id"], user.get("first_name") or "")


class PasswordLogin(BaseModel):
    login: str
    password: str


# Защита от перебора паролей: 5 ошибок по логину -> блокировка на 10 минут (в каждом воркере своя).
_FAILS: dict[str, tuple[int, float]] = {}


@router.post("/auth/password")
async def auth_password(body: PasswordLogin):
    import time as _time
    key = body.login.strip().lower()
    fails, until = _FAILS.get(key, (0, 0.0))
    if until > _time.time():
        raise HTTPException(429, "Ko'p marta xato kiritildi. 10 daqiqadan keyin urinib ko'ring.")
    row = await database.fetch_one(select(web_credentials).where(web_credentials.c.login == key))
    if not row or not verify_password(body.password, row["password_hash"]):
        fails += 1
        _FAILS[key] = (0, _time.time() + 600) if fails >= 5 else (fails, 0.0)
        raise HTTPException(401, "Login yoki parol noto'g'ri")
    _FAILS.pop(key, None)
    rights = await service.rights_of(row["telegram_user_id"])
    return await _login(row["telegram_user_id"], rights["name"])


class CredentialsChange(BaseModel):
    current_password: str
    new_login: Optional[str] = None
    new_password: Optional[str] = None


@router.get("/auth/credentials")
async def my_credentials(actor: dict = Depends(current_actor)):
    row = await database.fetch_one(select(web_credentials).where(web_credentials.c.telegram_user_id == actor["uid"]))
    return {"login": row["login"] if row else None}


@router.post("/auth/change-credentials")
async def change_credentials(body: CredentialsChange, actor: dict = Depends(current_actor)):
    row = await database.fetch_one(select(web_credentials).where(web_credentials.c.telegram_user_id == actor["uid"]))
    if not row:
        raise HTTPException(404, "Sizga hali login berilmagan")
    if not verify_password(body.current_password, row["password_hash"]):
        raise HTTPException(400, "Joriy parol noto'g'ri")
    vals = {}
    if body.new_login and body.new_login.strip().lower() != row["login"]:
        nl = body.new_login.strip().lower()
        if not (3 <= len(nl) <= 32) or not all(c.isascii() and (c.isalnum() or c in "._-") for c in nl):
            raise HTTPException(400, "Login: 3–32 belgi, faqat lotin harf, raqam, . _ -")
        if await database.fetch_one(select(web_credentials).where(web_credentials.c.login == nl)):
            raise HTTPException(400, "Bu login band")
        vals["login"] = nl
    if body.new_password:
        p = body.new_password
        if len(p) < 8 or p.isdigit() or p.isalpha():
            raise HTTPException(400, "Parol kamida 8 belgi, harf va raqam aralash bo'lsin")
        vals["password_hash"] = hash_password(p)
    if not vals:
        raise HTTPException(400, "O'zgarish yo'q")
    vals["updated_at"] = datetime.utcnow()
    await database.execute(web_credentials.update().where(web_credentials.c.telegram_user_id == actor["uid"]).values(**vals))
    await service.audit(actor["uid"], "change_credentials", "web_credentials", actor["uid"], {"login": row["login"]},
                        {"login": vals.get("login", row["login"]), "password_changed": "password_hash" in vals})
    return {"ok": True, "login": vals.get("login", row["login"])}


@router.get("/my")
async def my_attendance(from_: Optional[str] = Query(None, alias="from"), to: Optional[str] = None,
                        actor: dict = Depends(require_employee)):
    return await employee_card(actor["employee_no"], from_, to, actor={"level": "viewer"})


@router.get("/auth/dev")
async def auth_dev(uid: int):
    if not DEV_LOGIN:
        raise HTTPException(404)
    return await _login(uid, f"Dev {uid}")


@router.post("/auth/logout")
async def logout():
    resp = JSONResponse({"ok": True})
    resp.delete_cookie(COOKIE_NAME, path="/")
    return resp


@router.get("/me")
async def me(actor: dict = Depends(current_actor)):
    return actor


class DownloadLinkIn(BaseModel):
    path: str


@router.post("/download-link")
async def download_link(body: DownloadLinkIn, actor: dict = Depends(require_viewer)):
    path, _, query = body.path.partition("?")
    if not path.endswith(".xlsx") or not path.startswith("/"):
        raise HTTPException(400, "Faqat Excel fayllar uchun")
    full = router.prefix + path
    token = make_token(actor["uid"], actor["name"], only_path=full, ttl_seconds=120)
    sep = "&" if query else ""
    return {"url": f"{full}?{query}{sep}dl={token}"}


# ------------------------------ отчёты ------------------------------

@router.get("/employees")
async def list_employees(actor: dict = Depends(require_viewer)):
    ctx = await service.load_context()
    today = service.now_local().date()
    return [
        {"employee_no": no, "name": e["full_name"], "role": ctx.role_code_at(no, today),
         "role_name": ctx.role_names.get(ctx.role_code_at(no, today))}
        for no, e in sorted(ctx.employees.items(), key=lambda x: x[1]["full_name"])
        if e["active"] and not e["merged_into"]
    ]


@router.get("/day")
async def day_report(date_: Optional[str] = Query(None, alias="date"), role: Optional[str] = None, actor: dict = Depends(require_viewer)):
    d = _d(date_, service.now_local().date())
    ctx, shifts = await service.compute_range(d, d)
    rows = [service.shift_to_dict(ctx, s) for s in shifts if not role or s.role == role]
    return {
        "date": d.isoformat(),
        "rows": rows,
        "summary": {
            "people": len({r["employee_no"] for r in rows}),
            "by_status": dict(Counter(r["status"] for r in rows)),
            "by_role": dict(Counter(r["role_name"] for r in rows)),
        },
        "roles": ctx.role_names,
    }


@router.get("/employee/{employee_no}")
async def employee_card(employee_no: str, from_: Optional[str] = Query(None, alias="from"), to: Optional[str] = None,
                        actor: dict = Depends(require_viewer)):
    d_from, d_to = _range(from_, to)
    ctx = await service.load_context()
    emp = ctx.canonical(employee_no)
    if emp not in ctx.employees:
        raise HTTPException(404, "Xodim topilmadi")
    _, shifts = await service.compute_range(d_from, d_to, only=emp, ctx=ctx)
    rows = [service.shift_to_dict(ctx, s, detail=True) for s in shifts]
    aliases = [no for no, e in ctx.employees.items() if e.get("merged_into") == emp]
    return {
        "employee": {"employee_no": emp, "name": service.name_of(ctx, emp), "aliases": aliases,
                     "role": ctx.role_code_at(emp, d_to), "role_name": ctx.role_names.get(ctx.role_code_at(emp, d_to)),
                     "roles": [{"from": f, "to": t, "role": r} for f, t, r in ctx.role_periods.get(emp, [])]},
        "from": d_from.isoformat(), "to": d_to.isoformat(),
        "shifts": rows,
        "totals": {"span_sec": sum(r["span_sec"] or 0 for r in rows), "inside_sec": sum(r["inside_sec"] or 0 for r in rows),
                   "days": len(rows), "incomplete": sum(1 for r in rows if r["status"] == "to'liq emas")},
    }


@router.get("/now")
async def now_status(actor: dict = Depends(require_viewer)):
    """Текущее состояние: на ферме / работает снаружи / ушёл / сегодня не приходил."""
    from davomat.engine import work_date_of, debounce
    from datetime import time as _time
    now = service.now_local()
    today = work_date_of(now, _time(3, 0))
    day_start = datetime.combine(today, _time(3, 0))
    ctx = await service.load_context()
    punches = await service.load_punches(ctx, now - timedelta(days=14), now + timedelta(minutes=5))
    groups = {"inside": [], "outside_work": [], "left": [], "absent": []}
    for emp, info in ctx.employees.items():
        if not info["active"] or info.get("merged_into"):
            continue
        ps = punches.get(emp, [])
        if not ps:
            continue  # не появлялся 14 дней — чтобы не засорять список
        kept, _ = debounce(ps, 60)
        last = kept[-1]
        todays = [p for p in kept if p.ts >= day_start]
        row = {
            "employee_no": emp, "name": info["full_name"],
            "role_name": ctx.role_names.get(ctx.role_code_at(emp, today), ""),
            "last_ts": last.ts.strftime("%Y-%m-%d %H:%M:%S"), "last_direction": last.direction,
            "last_source": last.source,
            "first_today": todays[0].ts.strftime("%Y-%m-%d %H:%M:%S") if todays else None,
            "since_sec": int((now - last.ts).total_seconds()),
        }
        # Ночная смена охранника может продолжаться со вчерашнего дня — поэтому
        # «сейчас внутри» определяется по последней отметке (за 16 часов), независимо от границы дня.
        recent = (now - last.ts) <= timedelta(hours=16)
        if recent and last.direction == "in" and last.source == "telegram" and last.outside:
            groups["outside_work"].append(row)
        elif recent and last.direction == "in":
            groups["inside"].append(row)
        elif todays:
            groups["left"].append(row)
        else:
            groups["absent"].append(row)
    for g in groups.values():
        g.sort(key=lambda r: r["name"])
    return {"now": now.strftime("%Y-%m-%d %H:%M:%S"), "work_date": today.isoformat(),
            "counts": {k: len(v) for k, v in groups.items()}, "groups": groups}


@router.get("/guards")
async def guard_shifts(from_: Optional[str] = Query(None, alias="from"), to: Optional[str] = None, actor: dict = Depends(require_viewer)):
    d_from, d_to = _range(from_, to)
    ctx, shifts = await service.compute_range(d_from, d_to)
    rows = [service.shift_to_dict(ctx, s) for s in shifts if s.role == "qorovul"]
    have = {(r["work_date"], r["shift_type"]) for r in rows}
    today = service.now_local().date()
    empty = []
    x = d_from
    while x <= min(d_to, today):
        for t in (SHIFT_GUARD_DAY, SHIFT_GUARD_NIGHT):
            if (x.isoformat(), t) not in have:
                empty.append({"date": x.isoformat(), "shift_type": t})
        x += timedelta(days=1)
    guards = [no for no in ctx.employees if ctx.role_code_at(no, today) == "qorovul"]
    return {"from": d_from.isoformat(), "to": d_to.isoformat(), "rows": rows, "empty": empty, "guards_count": len(guards)}


def _period_rows(ctx, shifts, role: Optional[str]):
    agg: dict[str, dict] = defaultdict(lambda: {"days": 0, "closed": 0, "incomplete": 0, "span_sec": 0, "inside_sec": 0})
    for s in shifts:
        if role and s.role != role:
            continue
        a = agg[s.employee_no]
        a["days"] += 1
        a["closed"] += s.status in ("yopildi", "qo'lda tuzatilgan")
        a["incomplete"] += s.status == "to'liq emas"
        a["span_sec"] += s.span_sec or 0
        a["inside_sec"] += s.inside_sec or 0
        a["role_name"] = ctx.role_names.get(s.role, s.role)
    return [{"employee_no": no, "name": service.name_of(ctx, no), **v}
            for no, v in sorted(agg.items(), key=lambda x: service.name_of(ctx, x[0]))]


@router.get("/period")
async def period_report(from_: Optional[str] = Query(None, alias="from"), to: Optional[str] = None, role: Optional[str] = None,
                        actor: dict = Depends(require_viewer)):
    d_from, d_to = _range(from_, to, 30)
    ctx, shifts = await service.compute_range(d_from, d_to)
    return {"from": d_from.isoformat(), "to": d_to.isoformat(), "rows": _period_rows(ctx, shifts, role)}


def _xlsx(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for i, h in enumerate(header, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(h) + 4)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_response(data: bytes, name: str) -> StreamingResponse:
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


@router.get("/period.xlsx")
async def period_xlsx(from_: Optional[str] = Query(None, alias="from"), to: Optional[str] = None, role: Optional[str] = None,
                      actor: dict = Depends(require_viewer)):
    d_from, d_to = _range(from_, to, 30)
    ctx, shifts = await service.compute_range(d_from, d_to)
    summary = [[r["name"], r["role_name"], r["days"], r["closed"], r["incomplete"],
                fmt_duration(r["span_sec"]), fmt_duration(r["inside_sec"]),
                round(r["span_sec"] / 3600, 2), round(r["inside_sec"] / 3600, 2)] for r in _period_rows(ctx, shifts, role)]
    daily = []
    for s in shifts:
        if role and s.role != role:
            continue
        daily.append([s.work_date.isoformat(), service.name_of(ctx, s.employee_no), ctx.role_names.get(s.role, s.role),
                      s.shift_type, s.kirish.strftime("%d.%m %H:%M") if s.kirish else "", s.chiqish.strftime("%d.%m %H:%M") if s.chiqish else "",
                      fmt_duration(s.span_sec), fmt_duration(s.inside_sec), s.status, ", ".join(sorted(set(s.flags)))])
    data = _xlsx([
        ("Jami", ["F.I.O.", "Rol", "Kunlar", "Yopilgan", "To'liq emas", "Ish soati", "Fermada", "Ish soati (soat)", "Fermada (soat)"], summary),
        ("Kunlik", ["Sana", "F.I.O.", "Rol", "Smena", "Kirish", "Chiqish", "Ish soati", "Fermada", "Holat", "Belgilar"], daily),
    ])
    return _xlsx_response(data, f"davomat_{d_from}_{d_to}.xlsx")


async def _checkin_rows(d_from: date, d_to: date, employee_no: Optional[str] = None):
    start = datetime.combine(d_from, datetime.min.time())
    end = datetime.combine(d_to + timedelta(days=1), datetime.min.time())
    q = select(checkins).where(and_(checkins.c.ts >= start, checkins.c.ts < end))
    if employee_no:
        q = q.where(checkins.c.employee_no == employee_no)
    ctx = await service.load_context()
    rows = []
    for r in await database.fetch_all(q.order_by(checkins.c.ts.desc())):
        rows.append({"id": r["id"], "employee_no": r["employee_no"], "name": service.name_of(ctx, r["employee_no"]),
                     "direction": r["direction"], "ts": r["ts"].strftime("%Y-%m-%d %H:%M:%S"),
                     "lat": r["lat"], "lon": r["lon"], "accuracy": r["accuracy"],
                     "distance_m": r["distance_m"], "outside": bool(r["outside"])})
    return rows


@router.get("/checkins")
async def checkin_list(from_: Optional[str] = Query(None, alias="from"), to: Optional[str] = None, actor: dict = Depends(require_viewer)):
    d_from, d_to = _range(from_, to, 1)
    return {"from": d_from.isoformat(), "to": d_to.isoformat(), "rows": await _checkin_rows(d_from, d_to)}


@router.get("/checkins.xlsx")
async def checkin_xlsx(from_: Optional[str] = Query(None, alias="from"), to: Optional[str] = None, actor: dict = Depends(require_viewer)):
    d_from, d_to = _range(from_, to, 1)
    rows = await _checkin_rows(d_from, d_to)
    data = _xlsx([("Check-in", ["Vaqt", "F.I.O.", "Harakat", "Hududda", "Masofa (m)", "Aniqlik (m)", "Kenglik", "Uzunlik", "Xarita"],
                   [[r["ts"], r["name"], "Ishni boshladi" if r["direction"] == "in" else "Ishni tugatdi",
                     "Yo'q" if r["outside"] else "Ha", round(r["distance_m"]) if r["distance_m"] is not None else "",
                     round(r["accuracy"]) if r["accuracy"] is not None else "", r["lat"], r["lon"],
                     f"https://maps.google.com/?q={r['lat']},{r['lon']}"] for r in rows])])
    return _xlsx_response(data, f"checkin_{d_from}_{d_to}.xlsx")


@router.get("/checkins/{checkin_id}/photo")
async def checkin_photo(checkin_id: int, actor: dict = Depends(current_actor)):
    r = await database.fetch_one(select(checkins).where(checkins.c.id == checkin_id))
    if not r:
        raise HTTPException(404)
    if actor["level"] not in VIEW_LEVELS and actor["employee_no"] != r["employee_no"]:
        raise HTTPException(403)
    path = Path(r["photo_path"])
    if not path.is_file():
        raise HTTPException(404, "Rasm topilmadi")
    return FileResponse(path, media_type="image/jpeg", headers={"Cache-Control": "private, max-age=86400"})


# ------------------------------ check-in ------------------------------

def haversine_m(lat1, lon1, lat2, lon2) -> float:
    r = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = math.radians(lat2 - lat1), math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


class ClientLog(BaseModel):
    event: str
    detail: str = ""


_CLIENT_LOG_COUNT: dict[int, tuple[int, float]] = {}


@router.post("/client-log")
async def client_log(body: ClientLog, actor: dict = Depends(current_actor)):
    """Пишет в лог ошибки камеры/геопозиции с телефона для диагностики (не более 20 в минуту)."""
    import logging, time as _t
    n, since = _CLIENT_LOG_COUNT.get(actor["uid"], (0, _t.time()))
    if _t.time() - since > 60:
        n, since = 0, _t.time()
    if n >= 20:
        return {"ok": False}
    _CLIENT_LOG_COUNT[actor["uid"]] = (n + 1, since)
    logging.getLogger("davomat.client").warning("client uid=%s %s: %s", actor["uid"], body.event[:40], body.detail[:600])
    return {"ok": True}


@router.get("/checkin/state")
async def checkin_state(actor: dict = Depends(require_employee)):
    s = await service.get_settings()
    perms = await service.checkin_permissions(actor["employee_no"])
    return {"employee_name": actor["employee_name"], **perms,
            "chat_location": await service.fresh_chat_location(actor["uid"]),
            "max_accuracy_m": float(s["max_accuracy_m"]), "geofence": bool(s["farm_lat"] and s["farm_lon"])}


@router.post("/checkin")
async def checkin(direction: str = Form(...), lat: Optional[float] = Form(None), lon: Optional[float] = Form(None),
                  accuracy: Optional[float] = Form(None), loc_source: str = Form("device"),
                  photo: UploadFile = File(...), actor: dict = Depends(require_employee)):
    if direction not in ("in", "out"):
        raise HTTPException(400, "Harakat noto'g'ri")
    perms = await service.checkin_permissions(actor["employee_no"])
    if direction == "out" and not perms["can_out"]:
        raise HTTPException(409, "Avval «Ishni boshladim» ni belgilang — boshlanmagan ishni tugatib bo'lmaydi.")
    if direction == "in" and not perms["can_in"]:
        raise HTTPException(409, "Ish allaqachon boshlangan. Avval «Ishni tugatdim» ni belgilang.")
    if loc_source == "chat":
        # Координаты берутся не от клиента, а из геопозиции, отправленной в чат бота (GPS/трансляция).
        cl = await service.fresh_chat_location(actor["uid"])
        if not cl:
            raise HTTPException(400, "Chatdagi joylashuv topilmadi yoki eskirgan. Botga jonli joylashuvni qayta ulashing.")
        lat, lon, accuracy = cl["lat"], cl["lon"], cl["accuracy"] if cl["accuracy"] is not None else 1.0
    if lat is None or lon is None or not (-90 <= lat <= 90 and -180 <= lon <= 180):
        raise HTTPException(400, "Koordinata noto'g'ri")
    s = await service.get_settings()
    max_acc = float(s["max_accuracy_m"])
    if accuracy is None or accuracy <= 0 or accuracy > max_acc:
        raise HTTPException(400, f"Joylashuv aniqligi past ({accuracy or '?'} m). GPS yoqilganini tekshiring, "
                                 f"ochiq joyga chiqib qayta urinib ko'ring (kerak: {max_acc:.0f} m dan yaxshi).")
    data = await photo.read()
    if len(data) < 5_000 or len(data) > 8_000_000 or not data.startswith(b"\xff\xd8"):
        raise HTTPException(400, "Rasm noto'g'ri. Kamera orqali qayta suratga oling.")

    now = service.now_local()
    recent = await database.fetch_one(select(checkins).where(and_(
        checkins.c.employee_no == actor["employee_no"], checkins.c.direction == direction,
        checkins.c.ts >= now - timedelta(minutes=2))))
    if recent:
        raise HTTPException(409, "Siz hozirgina belgilagansiz.")

    distance = outside = None
    if s["farm_lat"] and s["farm_lon"]:
        distance = haversine_m(lat, lon, float(s["farm_lat"]), float(s["farm_lon"]))
        outside = distance > float(s["farm_radius_m"])

    folder = PHOTO_DIR / now.strftime("%Y/%m")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{now.strftime('%Y%m%d_%H%M%S')}_{actor['employee_no']}_{uuid.uuid4().hex[:8]}.jpg"
    path.write_bytes(data)

    cid = await database.execute(checkins.insert().values(
        employee_no=actor["employee_no"], telegram_user_id=actor["uid"], direction=direction, ts=now,
        lat=lat, lon=lon, accuracy=accuracy, distance_m=distance, outside=1 if outside else 0,
        photo_path=str(path), created_at=datetime.utcnow()))

    label = "▶️ Ishni boshladi" if direction == "in" else "⏹ Ishni tugatdi"
    where = ""
    if outside:
        where = f"\n📍 Ferma hududidan tashqarida ({distance / 1000:.1f} km)"
    elif distance is not None:
        where = "\n🏠 Ferma hududida"
    text = (f"{label}\n\n{_format_time_uz(now)}\n\n{html.escape(actor['employee_name'] or '')}{where}\n"
            f"<a href=\"https://maps.google.com/?q={lat},{lon}\">Xaritada ko'rish</a>")
    try:
        await telegram.notify(text, image_bytes=data, image_name=path.name)
    except Exception:
        pass
    return {"ok": True, "id": cid, "ts": now.strftime("%Y-%m-%d %H:%M:%S"),
            "outside": bool(outside), "distance_m": round(distance) if distance is not None else None}


# ------------------------------- админ -------------------------------

@router.get("/admin/employees")
async def admin_employees(actor: dict = Depends(require_admin)):
    ctx = await service.load_context()
    today = service.now_local().date()
    return [
        {**{k: e[k] for k in ("employee_no", "full_name", "telegram_user_id", "merged_into", "active")},
         "role": ctx.role_code_at(no, today), "role_auto": ctx.role_is_auto(no, today),
         "roles": [{"from": f, "to": t, "role": r} for f, t, r in ctx.role_periods.get(no, [])]}
        for no, e in sorted(ctx.employees.items(), key=lambda x: x[1]["full_name"])
    ]


class EmployeeUpdate(BaseModel):
    full_name: Optional[str] = None
    active: Optional[bool] = None
    merged_into: Optional[str] = None
    telegram_user_id: Optional[int] = None
    clear_telegram: bool = False
    clear_merge: bool = False


@router.post("/admin/employees/{employee_no}")
async def admin_update_employee(employee_no: str, body: EmployeeUpdate, actor: dict = Depends(require_admin)):
    before = await database.fetch_one(select(employees).where(employees.c.employee_no == employee_no))
    if not before:
        raise HTTPException(404, "Xodim topilmadi")
    vals = {}
    if body.full_name is not None and body.full_name.strip():
        vals["full_name"] = body.full_name.strip()
    if body.active is not None:
        vals["active"] = 1 if body.active else 0
    if body.clear_merge:
        vals["merged_into"] = None
    elif body.merged_into:
        if body.merged_into == employee_no:
            raise HTTPException(400, "O'ziga birlashtirib bo'lmaydi")
        if not await database.fetch_one(select(employees).where(employees.c.employee_no == body.merged_into)):
            raise HTTPException(400, "Asosiy xodim topilmadi")
        vals["merged_into"] = body.merged_into
    if body.clear_telegram:
        vals["telegram_user_id"] = None
    elif body.telegram_user_id:
        other = await database.fetch_one(select(employees).where(and_(
            employees.c.telegram_user_id == body.telegram_user_id, employees.c.employee_no != employee_no)))
        if other:
            raise HTTPException(400, f"Bu Telegram ID allaqachon {other['full_name']} ga biriktirilgan")
        vals["telegram_user_id"] = body.telegram_user_id
    if not vals:
        return {"ok": True}
    await database.execute(employees.update().where(employees.c.employee_no == employee_no).values(**vals))
    await service.audit(actor["uid"], "update", "employee", employee_no, dict(before), vals)
    old_uid, new_uid = before["telegram_user_id"], vals.get("telegram_user_id", before["telegram_user_id"])
    if old_uid and old_uid != new_uid:
        await service.drop_credentials_if_orphan(old_uid)
        if not (await service.rights_of(old_uid))["level"]:
            await service.notify_unlinked(old_uid)
    if new_uid and new_uid != old_uid:
        await database.execute(tg_link_requests.update().where(
            tg_link_requests.c.telegram_user_id == new_uid).values(status="linked"))
        await service.notify_linked(new_uid)
    return {"ok": True}


class RoleAssign(BaseModel):
    role: str
    valid_from: str


@router.post("/admin/employees/{employee_no}/role")
async def admin_assign_role(employee_no: str, body: RoleAssign, actor: dict = Depends(require_admin)):
    # «avto» — отменяет назначенную вручную роль с этой даты, система определяет сама.
    auto = body.role == "avto"
    if not auto and not await database.fetch_one(select(roles).where(roles.c.code == body.role)):
        raise HTTPException(400, "Rol topilmadi")
    vf = _d(body.valid_from, service.now_local().date())
    async with database.transaction():
        # Следующие периоды не удаляем — только закрываем открытый предыдущий период до этой даты.
        open_rows = await database.fetch_all(select(employee_roles).where(and_(
            employee_roles.c.employee_no == employee_no, employee_roles.c.valid_to.is_(None),
            employee_roles.c.valid_from < vf.isoformat())))
        for r in open_rows:
            await database.execute(employee_roles.update().where(employee_roles.c.id == r["id"])
                                   .values(valid_to=(vf - timedelta(days=1)).isoformat()))
        same_day = await database.fetch_all(select(employee_roles).where(and_(
            employee_roles.c.employee_no == employee_no, employee_roles.c.valid_from == vf.isoformat())))
        for r in same_day:
            await database.execute(employee_roles.delete().where(employee_roles.c.id == r["id"]))
        if not auto:
            await database.execute(employee_roles.insert().values(
                employee_no=employee_no, role=body.role, valid_from=vf.isoformat(), created_by=actor["uid"],
                created_at=datetime.utcnow()))
    service.reset_guard_cache()
    await service.audit(actor["uid"], "assign_role", "employee", employee_no,
                        [dict(r) for r in open_rows + same_day], {"role": body.role, "valid_from": vf.isoformat()})
    return {"ok": True}


@router.get("/admin/roles")
async def admin_roles(actor: dict = Depends(require_admin)):
    return [{"code": r["code"], "name": r["name"], "params": json.loads(r["params"])}
            for r in await database.fetch_all(select(roles))]


class RoleUpdate(BaseModel):
    name: Optional[str] = None
    params: dict


@router.put("/admin/roles/{code}")
async def admin_update_role(code: str, body: RoleUpdate, actor: dict = Depends(require_superadmin)):
    before = await database.fetch_one(select(roles).where(roles.c.code == code))
    if not before:
        raise HTTPException(404)
    try:
        RoleParams.from_dict(code, body.params)
    except Exception as e:
        raise HTTPException(400, f"Parametr noto'g'ri: {e}")
    vals = {"params": json.dumps(body.params)}
    if body.name:
        vals["name"] = body.name
    await database.execute(roles.update().where(roles.c.code == code).values(**vals))
    await service.audit(actor["uid"], "update", "role", code, dict(before), vals)
    return {"ok": True}


@router.get("/admin/settings")
async def admin_settings(actor: dict = Depends(require_admin)):
    return await service.get_settings()


@router.put("/admin/settings")
async def admin_update_settings(body: dict, actor: dict = Depends(require_superadmin)):
    before = await service.get_settings()
    changed = {}
    for k, v in body.items():
        if k not in service.SETTING_DEFAULTS:
            raise HTTPException(400, f"Noma'lum sozlama: {k}")
        v = str(v).strip()
        if k in ("farm_lat", "farm_lon", "farm_radius_m", "max_accuracy_m", "guard_grace_min",
                 "remind_after_h", "auto_resolve_after_h") and v:
            try:
                float(v)
            except ValueError:
                raise HTTPException(400, f"{k} son bo'lishi kerak")
        if before.get(k) != v:
            changed[k] = v
            if await database.fetch_one(select(settings).where(settings.c.key == k)):
                await database.execute(settings.update().where(settings.c.key == k).values(value=v))
            else:
                await database.execute(settings.insert().values(key=k, value=v))
    if changed:
        await service.audit(actor["uid"], "update", "settings", None, {k: before.get(k) for k in changed}, changed)
        service.reset_guard_cache()
    return {"ok": True, "changed": changed}


class CorrectionIn(BaseModel):
    employee_no: str
    work_date: str
    field: str
    value: Optional[str] = None  # «YYYY-MM-DD HH:MM» или пусто
    note: str = ""


@router.post("/admin/corrections")
async def admin_add_correction(body: CorrectionIn, actor: dict = Depends(require_admin)):
    if body.field not in ("kirish", "chiqish"):
        raise HTTPException(400, "Maydon noto'g'ri")
    wd = _d(body.work_date, service.now_local().date())
    val = None
    if body.value:
        try:
            val = datetime.strptime(body.value.strip()[:16], "%Y-%m-%d %H:%M")
        except ValueError:
            raise HTTPException(400, "Vaqt formati: YYYY-MM-DD HH:MM")
    if not body.note.strip():
        raise HTTPException(400, "Tuzatish sababini yozing")
    cid = await database.execute(corrections.insert().values(
        employee_no=body.employee_no, work_date=wd.isoformat(), field=body.field, value=val,
        note=body.note.strip(), active=1, created_by=actor["uid"], created_at=datetime.utcnow()))
    await service.audit(actor["uid"], "add", "correction", cid, None, body.model_dump())
    return {"ok": True, "id": cid}


@router.get("/admin/corrections")
async def admin_list_corrections(employee_no: Optional[str] = None, actor: dict = Depends(require_admin)):
    q = select(corrections).where(corrections.c.active == 1).order_by(corrections.c.id.desc()).limit(200)
    if employee_no:
        q = q.where(corrections.c.employee_no == employee_no)
    return [dict(r) for r in await database.fetch_all(q)]


@router.delete("/admin/corrections/{cid}")
async def admin_delete_correction(cid: int, actor: dict = Depends(require_admin)):
    before = await database.fetch_one(select(corrections).where(corrections.c.id == cid))
    if not before:
        raise HTTPException(404)
    await database.execute(corrections.update().where(corrections.c.id == cid).values(active=0))
    await service.audit(actor["uid"], "delete", "correction", cid, dict(before), {"active": 0})
    return {"ok": True}


class ResolutionIn(BaseModel):
    employee_no: str
    shift_start: str
    shift_type: str


@router.post("/admin/resolutions")
async def admin_resolve(body: ResolutionIn, actor: dict = Depends(require_admin)):
    await service.save_resolution(body.employee_no, datetime.fromisoformat(body.shift_start), body.shift_type, actor["uid"])
    return {"ok": True}


@router.get("/admin/users")
async def admin_users(actor: dict = Depends(require_superadmin)):
    logins = {r["telegram_user_id"]: r["login"] for r in await database.fetch_all(select(web_credentials))}
    return [{**dict(r), "login": logins.get(r["telegram_user_id"])}
            for r in await database.fetch_all(select(app_users).order_by(app_users.c.level, app_users.c.name))]


class UserIn(BaseModel):
    telegram_user_id: int
    name: str = ""
    level: str
    is_responsible: bool = False


def _can_manage(actor: dict, target_level) -> bool:
    """Суперадмин управляет всеми; админ — только viewer'ами (и теми, у кого ещё нет уровня)."""
    if actor["level"] == SUPERADMIN:
        return target_level != SUPERADMIN
    return target_level in (None, "viewer")


@router.post("/admin/users")
async def admin_upsert_user(body: UserIn, actor: dict = Depends(require_superadmin)):
    if body.level != "admin":
        raise HTTPException(400, "Faqat admin berish mumkin (super admin .env da belgilanadi)")
    before = await database.fetch_one(select(app_users).where(app_users.c.telegram_user_id == body.telegram_user_id))
    if before and before["level"] == SUPERADMIN:
        raise HTTPException(400, "Super admin .env dagi SUPERADMIN_TELEGRAM_ID orqali belgilanadi")
    if not _can_manage(actor, before["level"] if before else None) or (body.level == "admin" and actor["level"] != SUPERADMIN):
        raise HTTPException(403, "Adminlarni faqat super admin qo'sha va o'zgartira oladi")
    vals = {"name": body.name, "level": body.level, "is_responsible": 1 if body.is_responsible else 0}
    if before:
        await database.execute(app_users.update().where(app_users.c.telegram_user_id == body.telegram_user_id).values(**vals))
    else:
        await database.execute(app_users.insert().values(telegram_user_id=body.telegram_user_id, created_at=datetime.utcnow(), **vals))
    await service.audit(actor["uid"], "upsert", "app_user", body.telegram_user_id, dict(before) if before else None, vals)
    await database.execute(tg_link_requests.update().where(
        tg_link_requests.c.telegram_user_id == body.telegram_user_id).values(status="linked"))
    if not before or before["level"] != body.level:
        await service.notify_linked(body.telegram_user_id)
    return {"ok": True}


@router.delete("/admin/users/{uid}")
async def admin_delete_user(uid: int, actor: dict = Depends(require_superadmin)):
    before = await database.fetch_one(select(app_users).where(app_users.c.telegram_user_id == uid))
    if not before:
        raise HTTPException(404)
    if before["level"] == SUPERADMIN:
        raise HTTPException(400, "Super adminni o'chirib bo'lmaydi")
    if not _can_manage(actor, before["level"]):
        raise HTTPException(403, "Adminni faqat super admin o'chira oladi")
    await database.execute(app_users.delete().where(app_users.c.telegram_user_id == uid))
    await service.audit(actor["uid"], "delete", "app_user", uid, dict(before), None)
    await service.drop_credentials_if_orphan(uid)
    rights = await service.rights_of(uid)
    if rights["employee_no"]:
        await service.notify_linked(uid)
    else:
        await service.notify_unlinked(uid)
    return {"ok": True}


@router.post("/admin/users/{uid}/reset-password")
async def admin_reset_password(uid: int, actor: dict = Depends(require_admin)):
    rights = await service.rights_of(uid)
    if not rights["level"] and not rights["employee_no"]:
        raise HTTPException(404, "Bu Telegram ID hech kimga biriktirilmagan")
    if uid != actor["uid"] and rights["level"] in ADMIN_LEVELS and actor["level"] != SUPERADMIN:
        raise HTTPException(403, "Admin parolini faqat super admin yangilay oladi")
    await service.notify_linked(uid, reset=True)
    await service.audit(actor["uid"], "reset_password", "web_credentials", uid)
    return {"ok": True}


@router.get("/admin/link-requests")
async def admin_link_requests(actor: dict = Depends(require_admin)):
    return [dict(r) for r in await database.fetch_all(
        select(tg_link_requests).where(tg_link_requests.c.status == "pending").order_by(tg_link_requests.c.created_at.desc()))]


@router.post("/admin/link-requests/{uid}/reject")
async def admin_reject_link(uid: int, actor: dict = Depends(require_admin)):
    await database.execute(tg_link_requests.update().where(tg_link_requests.c.telegram_user_id == uid).values(status="rejected"))
    await service.audit(actor["uid"], "reject", "link_request", uid)
    return {"ok": True}


@router.get("/admin/audit")
async def admin_audit(limit: int = 100, actor: dict = Depends(require_admin)):
    rows = await database.fetch_all(select(audit_log).order_by(audit_log.c.id.desc()).limit(min(limit, 500)))
    return [dict(r) for r in rows]
